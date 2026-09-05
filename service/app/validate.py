"""Sanity-check a generated .docx before it is handed back to the caller."""
from __future__ import annotations

from docx import Document


def validate_docx(path: str) -> None:
    doc = Document(path)  # raises if the file is not a well-formed .docx
    if not doc.paragraphs and not doc.tables:
        raise ValueError("Generated document has no content - conversion likely failed")
